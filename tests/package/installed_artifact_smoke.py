"""Smoke-test an installed pyToCE artifact outside its checkout."""

import argparse
import inspect
import json
import os
import subprocess
import sys
from contextlib import closing
from importlib import metadata, resources
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

_PUBLIC = (
    "set_workspace",
    "download_data",
    "prepare_backgrounds",
    "run_mfa",
    "run_lca",
    "run_asr",
    "run_optimization",
)
_SIGNATURES = {
    "set_workspace": (
        "(workspace: str | pathlib.Path, *, refresh: bool = False) -> "
        "pytoce.set_workspace.set_workspace.WorkspaceReport"
    ),
    "download_data": (
        "(*, refresh: bool = False) -> pytoce.download_data.download_data.DownloadDataReport"
    ),
    "prepare_backgrounds": (
        "(*, background_scenario: str | collections.abc.Sequence[str] = "
        "('remind_ssp2_ndc', 'remind_ssp2_npi', 'remind_ssp2_pkbudg1000', "
        "'remind_ssp2_pkbudg650'), premise_key: str | None = None) -> "
        "pytoce.prepare_backgrounds.prepare_backgrounds.PrepareBackgroundsReport"
    ),
    "run_mfa": (
        "(*, foreground_scenario: str | collections.abc.Sequence[str] = "
        "'remind_mfa_default', refresh: bool = False) -> "
        "pytoce.run_mfa.run_mfa.RunMFAReport"
    ),
    "run_lca": (
        "(*, foreground_scenario: str | collections.abc.Sequence[str] = "
        "'remind_mfa_default', background_scenario: str | collections.abc.Sequence[str] = "
        "('remind_ssp2_ndc', 'remind_ssp2_npi', 'remind_ssp2_pkbudg1000', "
        "'remind_ssp2_pkbudg650'), lcia_method: str | collections.abc.Sequence[str] = "
        "('ef31', 'pb_lcia'), upstream_analysis_year: int | "
        "collections.abc.Sequence[int] = (2019, 2050, 2100), refresh: bool = False) -> "
        "pytoce.run_lca.run_lca.RunLCAReport"
    ),
    "run_asr": (
        "(*, allocated_cc_result: str, foreground_scenario: str | "
        "collections.abc.Sequence[str] = 'remind_mfa_default', background_scenario: str | "
        "collections.abc.Sequence[str] = ('remind_ssp2_ndc', 'remind_ssp2_npi', "
        "'remind_ssp2_pkbudg1000', 'remind_ssp2_pkbudg650'), lcia_method: str | "
        "collections.abc.Sequence[str] = ('ef31', 'pb_lcia'), allocated_cc_case: str | "
        "collections.abc.Sequence[str] | None = None, refresh: bool = False) -> "
        "pytoce.run_asr.run_asr.RunASRReport"
    ),
    "run_optimization": (
        "(*, foreground_scenario: str | collections.abc.Sequence[str], allocated_cc_result: "
        "str, background_scenario: str | collections.abc.Sequence[str] | None = None, "
        "optimization_lcia_method: str | collections.abc.Sequence[str] = 'ef31', "
        "lcia_method: str | collections.abc.Sequence[str] = ('ef31', 'pb_lcia'), "
        "optimization_block: Union[Literal['demand_lifetime_frontier', "
        "'circularity_feedstock_anchor', 'integrated_frontier'], "
        "collections.abc.Sequence[Literal['demand_lifetime_frontier', "
        "'circularity_feedstock_anchor', 'integrated_frontier']]] = "
        "('demand_lifetime_frontier', 'circularity_feedstock_anchor', "
        "'integrated_frontier'), allocated_cc_case: str | collections.abc.Sequence[str] | "
        "None = None, initial_grid_step_pct: int = 20, "
        "endpoint_refinement_steps_pct: collections.abc.Sequence[int] = (10, 5), "
        "candidate_workers: int = 1, refresh: bool = False) -> "
        "pytoce.run_optimization.run_optimization.RunOptimizationReport"
    ),
}

_TYPING_CONSUMER = """from pathlib import Path
from typing import TYPE_CHECKING, assert_type

if TYPE_CHECKING:
    from pytoce import (
        download_data, prepare_backgrounds, run_asr, run_lca, run_mfa,
        run_optimization, set_workspace,
    )
    from pytoce.download_data.download_data import DownloadDataReport
    from pytoce.prepare_backgrounds.prepare_backgrounds import PrepareBackgroundsReport
    from pytoce.run_asr.run_asr import RunASRReport
    from pytoce.run_lca.run_lca import RunLCAReport
    from pytoce.run_mfa.run_mfa import RunMFAReport
    from pytoce.run_optimization.run_optimization import RunOptimizationReport
    from pytoce.set_workspace.set_workspace import WorkspaceReport

    workspace = set_workspace("example")
    acquired = download_data()
    backgrounds = prepare_backgrounds()
    mfa = run_mfa()
    lca = run_lca()
    asr = run_asr(allocated_cc_result="example")
    optimization = run_optimization(
        foreground_scenario="example", allocated_cc_result="example",
        background_scenario=("remind_ssp2_npi", "remind_ssp2_pkbudg650"),
        initial_grid_step_pct=20,
        endpoint_refinement_steps_pct=(10, 5, 1),
        candidate_workers=3,
    )
    assert_type(workspace, WorkspaceReport)
    assert_type(acquired, DownloadDataReport)
    assert_type(backgrounds, PrepareBackgroundsReport)
    assert_type(mfa, RunMFAReport)
    assert_type(lca, RunLCAReport)
    assert_type(asr, RunASRReport)
    assert_type(optimization, RunOptimizationReport)
    assert_type(workspace.workspace, Path)
    assert_type(acquired.sources[0].artifact_path, Path)
    assert_type(backgrounds.results[0].background_scenario, str)
    assert_type(mfa.results[0].manifest_path, Path)
    assert_type(lca.results[0].annual_scores, Path)
    assert_type(asr.results[0].annual_asr, Path)
    assert_type(optimization.products[0].product_identity, str)
"""


def check_installed_typing() -> None:
    """Check the distributed report types."""

    assert resources.files("pytoce").joinpath("py.typed").is_file()
    environment = {key: value for key, value in os.environ.items() if key != "PYTHONPATH"}
    invalid = (
        "    bad_path: int = workspace.workspace\n"
        '    run_optimization(foreground_scenario="example", allocated_cc_result=1)\n'
    )
    first_invalid_line = len(_TYPING_CONSUMER.splitlines())
    expected_errors = [
        ("reportAssignmentType", first_invalid_line),
        ("reportArgumentType", first_invalid_line + 1),
    ]
    with TemporaryDirectory(prefix="pytoce-installed-typing-") as temporary:
        root = Path(temporary)
        consumer = root / "consumer.py"
        config = root / "pyrightconfig.json"
        config.write_text(
            json.dumps(
                {
                    "include": ["consumer.py"],
                    "pythonVersion": f"{sys.version_info.major}.{sys.version_info.minor}",
                    "typeCheckingMode": "strict",
                    "autoSearchPaths": False,
                    "extraPaths": [],
                    "useLibraryCodeForTypes": False,
                }
            ),
            encoding="utf-8",
        )
        for suffix, exit_code, expected in (("", 0, []), (invalid, 1, expected_errors)):
            consumer.write_text(_TYPING_CONSUMER + suffix, encoding="utf-8")
            result = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "pyright",
                    "--project",
                    str(config),
                    "--pythonpath",
                    sys.executable,
                    "--outputjson",
                ],
                cwd=root,
                env=environment,
                check=False,
                capture_output=True,
                text=True,
            )
            assert result.returncode == exit_code, (result.stdout, result.stderr)
            diagnostics = json.loads(result.stdout)["generalDiagnostics"]
            assert all(
                item["severity"] == "error" and Path(item["file"]) == consumer
                for item in diagnostics
            ), diagnostics
            actual = [(item["rule"], item["range"]["start"]["line"]) for item in diagnostics]
            assert sorted(actual) == sorted(expected), diagnostics


def smoke(workspace: Path) -> None:
    """Assert lazy exports, signatures, resources, and disposable installation."""

    import pytoce

    assert Path(pytoce.__file__).resolve().is_relative_to(Path(sys.prefix).resolve())
    assert pytoce.__all__ == list(_PUBLIC)
    assert metadata.version("pytoce") == pytoce.__version__ == "1.0.0"
    assert not any(name.startswith("pytoce.") for name in sys.modules if name != "pytoce")
    for name in _PUBLIC:
        function = getattr(pytoce, name)
        assert callable(function)
        assert str(inspect.signature(function)) == _SIGNATURES[name]

    report = pytoce.set_workspace(workspace)
    assert report.status in {"computed", "reused"}
    assert report.workspace == workspace.resolve()
    for path in (
        report.data_raw,
        report.data_processed,
        report.results,
        report.parameter_catalogue,
        report.default_input_parameters,
        report.default_annual_input_values,
        report.foreground_scenario_templates,
        report.foreground_scenarios,
        report.allocated_cc_templates,
        report.allocated_cc_results,
        report.methodological_notes,
        report.recommended_citations,
    ):
        assert path.exists(), path
    with closing(
        load_workbook(report.methodological_notes / "LCI_per_unit.xlsx", read_only=True)
    ) as workbook:
        assert {"Read me", "Providers", "Recipe quantities", "Sources"} <= set(workbook.sheetnames)
    check_installed_typing()


def main() -> None:
    """Validate one installed artifact and its disposable workspace."""

    parser = argparse.ArgumentParser()
    parser.add_argument("workspace", type=Path)
    arguments = parser.parse_args()
    smoke(arguments.workspace.resolve())


if __name__ == "__main__":
    main()
