"""Register default LCIA methods against the ecoinvent 3.12 biosphere.

Read the licensed EF 3.1 workbook and the ten selected PB-LCIA YAML members,
map their factors to biosphere flows, and register the selected methods.
"""

import json
import tarfile
from collections import defaultdict
from collections.abc import Iterable
from contextlib import closing
from dataclasses import dataclass
from hashlib import sha256
from importlib import resources
from math import isfinite
from pathlib import Path, PurePosixPath
from typing import Any, BinaryIO, Final, cast

import yaml
from openpyxl import load_workbook  # type: ignore[import-untyped]

from pytoce.shared.runtime.filesystem import canonical_json_bytes, payload_sha256

LCIA_CATALOGUE_RESOURCE: Final = "lcia_catalogue.json"


@dataclass(frozen=True, slots=True)
class MethodSpec:
    """One ordered default impact-category authority."""

    order: int
    method_id: str
    method_version: str
    impact_category_id: str
    impact_category_tuple: tuple[str, ...]
    unit: str
    source_member: str


@dataclass(frozen=True, slots=True)
class ResolvedMethod:
    """One exact method and its resolved Brightway factor data."""

    spec: MethodSpec
    factors: tuple[tuple[int, float], ...]
    factor_identity: str


def load_lcia_catalogue() -> tuple[tuple[MethodSpec, ...], dict[str, object]]:
    """Load the package-owned ordered impact-category catalog."""

    raw = json.loads(
        resources.files("pytoce.prepare_backgrounds.data")
        .joinpath(LCIA_CATALOGUE_RESOURCE)
        .read_text(encoding="utf-8")
    )
    catalogue_rows = cast(list[dict[str, object]], raw["impact_categories"])
    specs = tuple(
        sorted(
            (
                MethodSpec(
                    order=int(cast(int, row["order"])),
                    method_id=str(row["method_id"]),
                    method_version=str(row["method_version"]),
                    impact_category_id=str(row["impact_category_id"]),
                    impact_category_tuple=tuple(
                        str(value) for value in cast(list[object], row["tuple"])
                    ),
                    unit=str(row["unit"]),
                    source_member=str(row["source_member"]),
                )
                for row in catalogue_rows
            ),
            key=lambda spec: spec.order,
        )
    )
    return specs, raw


def _biosphere_index(
    database_name: str,
) -> tuple[
    dict[tuple[str, tuple[str, ...]], tuple[Any, ...]],
    dict[tuple[str, tuple[str, ...]], tuple[Any, ...]],
]:
    """Index biosphere flows by name and compartments for exact and casefold matching."""
    import bw2data as bd  # type: ignore[import-untyped]

    by_identity: dict[tuple[str, tuple[str, ...]], list[Any]] = defaultdict(list)
    by_casefold_identity: dict[tuple[str, tuple[str, ...]], list[Any]] = defaultdict(list)
    for flow in cast(Iterable[Any], bd.Database(database_name)):
        name = str(flow.get("name", ""))
        categories = tuple(str(value) for value in flow.get("categories", ()))
        by_identity[(name, categories)].append(flow)
        by_casefold_identity[
            (name.casefold(), tuple(value.casefold() for value in categories))
        ].append(flow)
    return (
        {key: tuple(value) for key, value in by_identity.items()},
        {key: tuple(value) for key, value in by_casefold_identity.items()},
    )


def _factor_identity(
    method_tuple: tuple[str, ...],
    factors: list[tuple[Any, float]],
) -> str:
    """Hash an ordered set of biosphere flow factors for one LCIA method."""
    payload = {
        "impact_category_tuple": list(method_tuple),
        "factors": sorted(
            (
                {
                    "database": str(flow["database"]),
                    "code": str(flow["code"]),
                    "amount": float(amount),
                }
                for flow, amount in factors
            ),
            key=lambda row: (row["database"], row["code"], row["amount"]),
        ),
    }
    return payload_sha256(payload)


def _resolved_method(
    spec: MethodSpec,
    factors: list[tuple[Any, float]],
) -> ResolvedMethod:
    """Resolve one catalog entry to Brightway biosphere-flow factors."""
    identity = _factor_identity(spec.impact_category_tuple, factors)
    return ResolvedMethod(
        spec=spec,
        factors=tuple(
            sorted(
                ((int(flow.id), float(amount)) for flow, amount in factors),
                key=lambda item: item[0],
            )
        ),
        factor_identity=identity,
    )


def _biosphere_categories(compartment: object, subcompartment: object) -> tuple[str, ...]:
    """Omit unspecified subcompartments from the ordered Brightway category tuple."""
    categories: tuple[str, ...] = (str(compartment),)
    if subcompartment not in (None, "", "unspecified"):
        categories += (str(subcompartment),)
    return categories


def resolve_ef_methods(
    workbook: Path,
    *,
    biosphere_database: str,
    specs: tuple[MethodSpec, ...],
) -> tuple[ResolvedMethod, ...]:
    """Resolve the selected EF v3.1 factors to ecoinvent 3.12 biosphere flows.

    Args:
        workbook: EF v3.1 workbook admitted by the download owner.
        biosphere_database: Name of the imported ecoinvent 3.12 biosphere database.
        specs: Ordered pyToCE impact-category definitions.

    Returns:
        EF methods in catalog order. Each factor retains its source amount and
        references the integer identifier of one exact Brightway biosphere flow.

    Raises:
        ValueError: If the workbook categories or units differ from the catalog, a
            factor does not resolve to one biosphere flow, or a factor is not finite.
    """

    expected = {spec.impact_category_tuple: spec for spec in specs if spec.method_id == "ef31"}
    with (
        workbook.open("rb") as stream,
        closing(load_workbook(stream, read_only=True, data_only=True)) as book,
    ):
        indicators: dict[tuple[str, ...], str] = {}
        for method, category, indicator, unit in book["Indicators"].iter_rows(
            min_row=2,
            values_only=True,
        ):
            method_tuple = (str(method), str(category), str(indicator))
            if method_tuple in expected:
                indicators[method_tuple] = str(unit)
        if set(indicators) != set(expected):
            raise ValueError("The EF workbook impact-category tuples changed")
        for indicator_tuple, unit in indicators.items():
            if unit != expected[indicator_tuple].unit:
                raise ValueError(f"EF unit changed for {indicator_tuple!r}")

        index, _casefold_index = _biosphere_index(biosphere_database)
        grouped: dict[tuple[str, ...], list[tuple[Any, float]]] = defaultdict(list)
        for row in book["CFs"].iter_rows(min_row=2, values_only=True):
            method, category, indicator, name, compartment, subcompartment, amount = row
            cf_method_tuple = (str(method), str(category), str(indicator))
            if cf_method_tuple not in expected:
                continue
            categories = _biosphere_categories(compartment, subcompartment)
            matches = index.get((str(name), categories), ())
            if len(matches) != 1:
                raise ValueError(
                    "EF factor does not resolve exactly: "
                    f"{name!r}, {categories!r}, matches={len(matches)}"
                )
            if (
                isinstance(amount, bool)
                or not isinstance(amount, (int, float))
                or not isfinite(float(amount))
            ):
                raise ValueError(
                    f"The EF factor {amount!r} for {cf_method_tuple!r} in {workbook} must be "
                    "a finite number. Leave this workspace unchanged. Place the documented "
                    "ecoinvent 3.12 files in a new workspace and run the workflows there."
                )
            grouped[cf_method_tuple].append((matches[0], float(amount)))
    return tuple(
        _resolved_method(expected[method_tuple], grouped[method_tuple])
        for method_tuple in sorted(expected, key=lambda value: expected[value].order)
    )


def _read_pb_members(archive: Path, selected: set[str]) -> dict[str, list[dict[str, object]]]:
    """Parse only selected PB factor tables from the admitted archive."""
    members: dict[str, list[dict[str, object]]] = {}
    with tarfile.open(archive, mode="r:gz") as handle:
        for member in handle.getmembers():
            name = PurePosixPath(member.name).name
            if not member.isfile() or name not in selected:
                continue
            with cast(BinaryIO, handle.extractfile(member)) as source:
                raw = yaml.safe_load(source.read())
            if not isinstance(raw, list) or any(not isinstance(row, dict) for row in raw):
                raise ValueError(f"PB-LCIA member {member.name!r} is not a factor table")
            members[name] = raw
    return members


def resolve_pb_methods(
    archive: Path,
    *,
    biosphere_database: str,
    specs: tuple[MethodSpec, ...],
) -> tuple[ResolvedMethod, ...]:
    """Resolve the ten selected PB-LCIA methods to ecoinvent 3.12 flows.

    Args:
        archive: PB-LCIA archive admitted by the download owner.
        biosphere_database: Name of the imported ecoinvent 3.12 biosphere database.
        specs: Ordered pyToCE impact-category definitions.

    Returns:
        The ten PB-LCIA methods in catalog order, with factors linked to
        Brightway biosphere-flow identifiers.

    Raises:
        ValueError: If required archive members are missing, a source flow is not an
            exact or case-insensitive unique match, or a factor is not finite.

    Notes:
        Case-insensitive matching is used only when the exact flow name and compartment
        do not resolve and the ecoinvent biosphere contains one unambiguous match.
    """

    expected = tuple(spec for spec in specs if spec.method_id == "pb_lcia")
    required_members = {spec.source_member for spec in expected}
    members = _read_pb_members(archive, required_members)
    if not required_members.issubset(members):
        raise ValueError("The PB-LCIA archive member set changed")
    index, casefold_index = _biosphere_index(biosphere_database)
    resolved: list[ResolvedMethod] = []
    for spec in expected:
        factors: list[tuple[Any, float]] = []
        for row in members[spec.source_member]:
            source_name = str(row["name"])
            categories = tuple(str(row["categories"]).split("::"))
            matches = index.get((source_name, categories), ())
            if len(matches) == 1:
                flow = matches[0]
            else:
                casefold_matches = casefold_index.get(
                    (
                        source_name.casefold(),
                        tuple(value.casefold() for value in categories),
                    ),
                    (),
                )
                if len(casefold_matches) != 1:
                    raise ValueError(
                        "PB factor does not resolve by an exact or casefold-unique "
                        "biosphere identity: "
                        f"{source_name!r}, {categories!r}, "
                        f"exact_matches={len(matches)}, "
                        f"casefold_matches={len(casefold_matches)}"
                    )
                flow = casefold_matches[0]
            amount = row["amount"]
            if (
                isinstance(amount, bool)
                or not isinstance(amount, (int, float))
                or not isfinite(float(amount))
            ):
                raise ValueError(
                    f"The PB-LCIA factor {amount!r} for {spec.impact_category_tuple!r} "
                    f"in {archive} must be a finite number. Leave this workspace unchanged. "
                    "Create a new workspace and run download_data before preparing the "
                    "backgrounds there."
                )
            factors.append((flow, float(amount)))
        resolved.append(_resolved_method(spec, factors))
    return tuple(resolved)


def register_methods(
    resolved: tuple[ResolvedMethod, ...],
    *,
    source_identity: str,
) -> tuple[dict[str, object], ...]:
    """Register the selected methods in the collision-free project."""

    import bw2data as bd  # type: ignore[import-untyped]

    for item in resolved:
        method = bd.Method(item.spec.impact_category_tuple)
        method.register(
            unit=item.spec.unit,
            source=item.spec.source_member,
            pytoce_method_id=item.spec.method_id,
            pytoce_method_version=item.spec.method_version,
            pytoce_factor_identity=item.factor_identity,
            pytoce_source_identity=source_identity,
            principal=True,
        )
        method.write(list(item.factors), process=True)
    return method_authority_rows(resolved)


def method_authority_rows(
    resolved: tuple[ResolvedMethod, ...],
) -> tuple[dict[str, object], ...]:
    """Return canonical scientific rows for resolved LCIA methods."""

    return tuple(
        {
            "selection_order": item.spec.order,
            "method_id": item.spec.method_id,
            "method_version": item.spec.method_version,
            "impact_category_id": item.spec.impact_category_id,
            "impact_category_tuple_json": json.dumps(
                item.spec.impact_category_tuple,
                ensure_ascii=True,
                separators=(",", ":"),
            ),
            "unit": item.spec.unit,
            "factor_identity": item.factor_identity,
        }
        for item in resolved
    )


def lcia_authority_identity(rows: tuple[dict[str, object], ...]) -> str:
    """Return the ordered 26-category registered-authority identity."""

    return "pytoce_lcia_authority_v4::" + sha256(canonical_json_bytes(list(rows))).hexdigest()
